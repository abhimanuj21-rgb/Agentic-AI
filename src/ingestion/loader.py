"""Local file loader with incremental hash-based change detection."""

import hashlib
import json
import os
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Set, Tuple


@dataclass
class LoadedDocument:
    """Represents a loaded file with its raw content and rich metadata."""
    content: str
    metadata: Dict[str, str] = field(default_factory=dict)


# Directories that should always be ignored when scanning recursively
IGNORE_DIRS: Set[str] = {
    ".git",
    ".venv",
    "venv",
    "env",
    "__pycache__",
    "node_modules",
    ".idea",
    ".vscode",
    "chroma_db",
    "dist",
    "build",
    ".cache",
}

# Default supported extensions for the coding agent knowledge base
DEFAULT_SUPPORTED_EXTENSIONS: Set[str] = {
    ".py",
    ".js",
    ".ts",
    ".jsx",
    ".tsx",
    ".html",
    ".css",
    ".md",
    ".txt",
    ".json",
    ".yaml",
    ".yml",
    ".pdf",
    ".docx",
}


def compute_file_hash(file_path: Path) -> str:
    """Calculate the SHA-256 hash of a file's contents."""
    sha256 = hashlib.sha256()
    with open(file_path, "rb") as f:
        while chunk := f.read(65536):
            sha256.update(chunk)
    return sha256.hexdigest()


def load_hash_manifest(manifest_path: str) -> Dict[str, Dict]:
    """Load the persisted incremental indexing hash manifest."""
    path = Path(manifest_path)
    if path.exists():
        try:
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return {}
    return {}


def save_hash_manifest(manifest_path: str, manifest: Dict[str, Dict]) -> None:
    """Save the updated incremental indexing hash manifest."""
    path = Path(manifest_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(manifest, f, indent=2, default=str)


def extract_text_from_file(file_path: Path) -> str:
    """Extract string content from various supported file types."""
    ext = file_path.suffix.lower()

    if ext in {".pdf"}:
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(file_path))
            pages_text = []
            for i, page in enumerate(reader.pages):
                extracted = page.extract_text()
                if extracted:
                    pages_text.append(f"--- Page {i + 1} ---\n{extracted}")
            return "\n\n".join(pages_text)
        except Exception as e:
            return f"[Error reading PDF file {file_path.name}: {str(e)}]"

    elif ext in {".docx"}:
        try:
            import docx
            doc = docx.Document(str(file_path))
            return "\n".join([p.text for p in doc.paragraphs if p.text])
        except Exception as e:
            return f"[Error reading DOCX file {file_path.name}: {str(e)}]"

    else:
        # Standard text and code files (.py, .md, .txt, .js, .ts, .json, .yaml, etc.)
        encodings = ["utf-8", "latin-1", "cp1252"]
        for enc in encodings:
            try:
                with open(file_path, "r", encoding=enc) as f:
                    return f.read()
            except UnicodeDecodeError:
                continue
            except Exception as e:
                return f"[Error reading file {file_path.name}: {str(e)}]"
        return f"[Error: Unable to decode file {file_path.name}]"


class LocalDirectoryLoader:
    """Scans and loads local files with support for incremental hash checks."""

    def __init__(
        self,
        base_path: str,
        supported_extensions: Optional[Set[str]] = None,
        manifest_path: str = "./data/chroma_db/file_hashes.json",
    ):
        self.base_path = Path(base_path).resolve()
        self.supported_extensions = supported_extensions or DEFAULT_SUPPORTED_EXTENSIONS
        self.manifest_path = manifest_path
        self.manifest = load_hash_manifest(self.manifest_path)

    def scan_files(self) -> List[Path]:
        """Recursively scan base_path for eligible files, ignoring unwanted directories."""
        if not self.base_path.exists():
            return []

        matched_files = []
        for root, dirs, files in os.walk(self.base_path):
            # Modify dirs in-place to prevent walking into ignored directories
            dirs[:] = [d for d in dirs if d not in IGNORE_DIRS and not d.startswith(".")]

            for file in files:
                file_path = Path(root) / file
                if file_path.suffix.lower() in self.supported_extensions:
                    matched_files.append(file_path)

        return matched_files

    def load_incremental(
        self,
    ) -> Tuple[List[LoadedDocument], List[str], Dict[str, Dict]]:
        """
        Load documents incrementally using SHA-256 hash checks.

        Returns:
            Tuple of:
            - List of LoadedDocument instances that are NEW or MODIFIED
            - List of file_path strings that were DELETED
            - Updated hash manifest dictionary
        """
        all_files = self.scan_files()
        new_or_modified_docs: List[LoadedDocument] = []
        current_file_paths: Set[str] = set()
        updated_manifest = dict(self.manifest)

        for file_path in all_files:
            rel_path_str = str(file_path.relative_to(self.base_path)).replace("\\", "/")
            current_file_paths.add(rel_path_str)

            file_hash = compute_file_hash(file_path)
            cached_info = self.manifest.get(rel_path_str)

            # Check if file has changed
            if cached_info and cached_info.get("hash") == file_hash:
                continue  # Unchanged, skip processing

            # File is new or modified
            content = extract_text_from_file(file_path)
            stat = file_path.stat()
            metadata = {
                "file_path": rel_path_str,
                "abs_path": str(file_path),
                "filename": file_path.name,
                "extension": file_path.suffix.lower(),
                "file_size": stat.st_size,
                "last_modified": datetime.fromtimestamp(stat.st_mtime).isoformat(),
                "hash": file_hash,
            }

            new_or_modified_docs.append(LoadedDocument(content=content, metadata=metadata))
            updated_manifest[rel_path_str] = {
                "hash": file_hash,
                "last_modified": metadata["last_modified"],
                "file_size": stat.st_size,
            }

        # Check for deleted files that are in manifest but no longer on disk
        deleted_files = [p for p in self.manifest.keys() if p not in current_file_paths]
        for deleted_path in deleted_files:
            updated_manifest.pop(deleted_path, None)

        return new_or_modified_docs, deleted_files, updated_manifest
